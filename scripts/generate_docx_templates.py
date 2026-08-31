import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_footer(doc, title):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run(f"HỆ THỐNG SOẠN ĐỀ THI | {title}")
    hrun.font.name = 'Arial'
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor(120, 120, 120)

def add_styled_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.bold = True
    if level == 1:
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(15, 23, 42)
    elif level == 2:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(30, 58, 138)
    return p

def add_tag(p, tag_name, tag_value=None, tag_color="0F172A", val_color="0284C7"):
    r_tag = p.add_run(tag_name + (" " if tag_value is not None else ""))
    r_tag.font.name = 'Arial'
    r_tag.font.size = Pt(10)
    r_tag.bold = True
    r_tag.font.color.rgb = RGBColor.from_string(tag_color)

    if tag_value is not None:
        r_val = p.add_run(str(tag_value))
        r_val.font.name = 'Arial'
        r_val.font.size = Pt(10)
        r_val.font.color.rgb = RGBColor.from_string(val_color)

def generate_exam_docx(output_path):
    doc = docx.Document()
    add_header_footer(doc, "FILE_DE_THI_DANH_CHO_GIAO_VIEN")

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    r_t = p_title.add_run("MẪU SOẠN ĐỀ THI (DÀNH CHO GIÁO VIÊN)")
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(16)
    r_t.bold = True
    r_t.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Tách biệt rõ ràng [NỘI DUNG CÂU HỎI] và [PHƯƠNG ÁN A/B/C/D] để Regex đọc chính xác 100%")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # 1. METADATA
    add_styled_heading(doc, "I. THÔNG TIN BÀI THI", level=1)
    
    meta_table = doc.add_table(rows=1, cols=1)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = meta_table.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

    p_meta = cell.paragraphs[0]
    p_meta.paragraph_format.space_after = Pt(4)
    add_tag(p_meta, "[BÀI THI]:", "Đọc hiểu & Ngôn ngữ - Module 1", tag_color="1E3A8A")
    
    p_pos = cell.add_paragraph()
    add_tag(p_pos, "[VỊ TRÍ BẮT ĐẦU]:", "1 (Nhập từ câu số 1)", tag_color="0F172A", val_color="0369A1")

    # 2. PASSAGE
    add_styled_heading(doc, "II. BÀI ĐỌC / NGỮ LIỆU (NẾU CÓ)", level=1)
    
    p_pass_intro = doc.add_paragraph()
    p_pass_intro.paragraph_format.space_after = Pt(4)
    add_tag(p_pass_intro, "[BÀI ĐỌC 1]", tag_color="047857")

    p_meta_pass = [
        ("[THỂ LOẠI]:", "Khoa học"),
        ("[NGUỒN TRÍCH]:", "Tác động của hạt vi nhựa đến hệ sinh thái biển"),
        ("[TÁC GIẢ]:", "Dr. Jane Smith"),
        ("[NĂM XUẤT BẢN]:", "2024")
    ]
    for tag, val in p_meta_pass:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_tag(p, tag, val, tag_color="065F46", val_color="0369A1")

    p_pbody = doc.add_paragraph()
    p_pbody.paragraph_format.space_before = Pt(4)
    p_pbody.paragraph_format.space_after = Pt(6)
    add_tag(p_pbody, "[NỘI DUNG BÀI ĐỌC]:", tag_color="065F46")
    r_body = p_pbody.add_run("\nRecent studies have highlighted the alarming prevalence of microplastics in aquatic environments. These minute synthetic particles, measuring less than five millimeters in diameter, originate from industrial discharge and degrading consumer waste. Marine organisms often mistake these particles for food, leading to severe bioaccumulation along the food chain. [Media:q01_marine_map.png]")
    r_body.font.name = 'Arial'
    r_body.font.size = Pt(10)

    # 3. QUESTIONS
    add_styled_heading(doc, "III. DANH SÁCH CÂU HỎI", level=1)

    # Q1: Trắc nghiệm (Tách nội dung và phương án)
    add_styled_heading(doc, "Câu 1: Câu hỏi Trắc nghiệm 4 phương án (Có thẻ tách biệt)", level=2)
    q1_p = doc.add_paragraph()
    add_tag(q1_p, "[CÂU 1]", tag_color="B91C1C")
    
    q1_tags = [
        ("[BÀI ĐỌC]:", "Bài đọc 1"),
        ("[DẠNG CÂU]:", "Trắc nghiệm"),
        ("[MỨC ĐỘ]:", "Trung bình"),
        ("[LĨNH VỰC]:", "Cấu trúc & Từ vựng"),
        ("[DẠNG BÀI]:", "Từ trong ngữ cảnh"),
        ("[THỜI GIAN LÀM]:", "75 giây")
    ]
    for tag, val in q1_tags:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_tag(p, tag, val, tag_color="6B21A8", val_color="0369A1")

    p_stem1 = doc.add_paragraph()
    p_stem1.paragraph_format.space_before = Pt(4)
    p_stem1.paragraph_format.space_after = Pt(4)
    add_tag(p_stem1, "[NỘI DUNG CÂU HỎI]:", tag_color="1E293B")
    r_s1 = p_stem1.add_run(" Which choice best completes the sentence with the most logical and appropriate word?")
    r_s1.font.name = 'Arial'
    r_s1.font.size = Pt(10)

    # ĐÁP ÁN ĐƯỢC TÁCH RIÊNG THÀNH CÁC THẺ [PHƯƠNG ÁN A], [PHƯƠNG ÁN B]...
    opts = [
        ("[PHƯƠNG ÁN A]:", "ubiquitous"),
        ("[PHƯƠNG ÁN B]:", "fleeting"),
        ("[PHƯƠNG ÁN C]:", "negligible"),
        ("[PHƯƠNG ÁN D]:", "extraneous")
    ]
    for tag, val in opts:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(2)
        add_tag(p, tag, val, tag_color="047857", val_color="0F172A")

    # Q2: Điền số
    add_styled_heading(doc, "Câu 2: Câu hỏi Điền số (Tự luận ngắn / Toán)", level=2)
    q2_p = doc.add_paragraph()
    add_tag(q2_p, "[CÂU 2]", tag_color="B91C1C")
    
    q2_tags = [
        ("[DẠNG CÂU]:", "Điền số"),
        ("[MỨC ĐỘ]:", "Khó"),
        ("[LĨNH VỰC]:", "Đại số"),
        ("[DẠNG BÀI]:", "Phương trình bậc nhất một ẩn"),
        ("[MÁY TÍNH]:", "Được dùng"),
        ("[THỜI GIAN LÀM]:", "90 giây"),
        ("[GỢI Ý]:", "Biến đổi biểu thức 6x + 3 theo cụm 3x + 9")
    ]
    for tag, val in q2_tags:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_tag(p, tag, val, tag_color="6B21A8", val_color="0369A1")

    p_stem2 = doc.add_paragraph()
    p_stem2.paragraph_format.space_before = Pt(4)
    p_stem2.paragraph_format.space_after = Pt(4)
    add_tag(p_stem2, "[NỘI DUNG CÂU HỎI]:", tag_color="1E293B")
    r_s2 = p_stem2.add_run(" If 3x + 9 = 21, what is the value of 6x + 3? [Media:q02_graph.png]")
    r_s2.font.name = 'Arial'
    r_s2.font.size = Pt(10)

    doc.save(output_path)

def generate_solution_docx(output_path):
    doc = docx.Document()
    add_header_footer(doc, "FILE_LOI_GIAI_DANH_CHO_GIAO_VIEN")

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    r_t = p_title.add_run("MẪU FILE LỜI GIẢI & PHÂN TÍCH ĐỀ THI")
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(16)
    r_t.bold = True
    r_t.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Khớp câu hỏi theo [CÂU 1], [CÂU 2]... - Dành cho Giáo viên soạn đáp án & giải thích")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # 1. METADATA
    add_styled_heading(doc, "I. THÔNG TIN BÀI THI KHỚP ĐỀ", level=1)
    
    meta_table = doc.add_table(rows=1, cols=1)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = meta_table.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

    p_meta = cell.paragraphs[0]
    add_tag(p_meta, "[BÀI THI]:", "Đọc hiểu & Ngôn ngữ - Module 1", tag_color="1E3A8A")

    # 2. QUICK ANSWER MATRIX
    add_styled_heading(doc, "II. BẢNG ĐÁP ÁN NHANH TRA CỨU", level=1)
    
    p_qans = doc.add_paragraph()
    p_qans.paragraph_format.space_after = Pt(4)
    add_tag(p_qans, "[BẢNG ĐÁP ÁN]", tag_color="047857")

    answers = [
        ("Câu 1", "A"),
        ("Câu 2", "27 (hoặc 54/2)")
    ]
    for q_code, ans in answers:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{q_code}: {ans}")
        r.font.name = 'Arial'
        r.font.size = Pt(10)
        r.bold = True
        r.font.color.rgb = RGBColor(30, 41, 59)

    # 3. DETAILED SOLUTIONS
    add_styled_heading(doc, "III. LỜI GIẢI CHI TIẾT & PHÂN TÍCH TỪNG CÂU", level=1)

    # Q1 Solution
    add_styled_heading(doc, "Lời giải Câu 1 (Trắc nghiệm)", level=2)
    doc_p1 = doc.add_paragraph()
    add_tag(doc_p1, "[CÂU 1]", tag_color="B91C1C")

    tags_q1 = [
        ("[ĐÁP ÁN ĐÚNG]:", "A"),
        ("[THÔNG SỐ IRT]:", "a=1.20, b=0.30, c=0.25 (Phân hóa=1.2, Độ khó=0.3, Đoán mò=0.25)")
    ]
    for tag, val in tags_q1:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_tag(p, tag, val, tag_color="6B21A8", val_color="0369A1")

    p_exp1 = doc.add_paragraph()
    p_exp1.paragraph_format.space_before = Pt(4)
    p_exp1.paragraph_format.space_after = Pt(4)
    add_tag(p_exp1, "[LỜI GIẢI CHI TIẾT]:", tag_color="047857")
    r_exp1 = p_exp1.add_run("\nTừ 'ubiquitous' (xuất hiện ở khắp mọi nơi) phù hợp nhất với ngữ cảnh 'alarming prevalence' (sự xuất hiện phổ biến đến mức báo động) của các hạt vi nhựa trong môi trường nước được mô tả ở bài đọc.")
    r_exp1.font.name = 'Arial'

    rationales_q1 = [
        ("[GIẢI THÍCH A]:", "ĐÚNG. Ubiquitous thể hiện sự hiện diện ở khắp mọi nơi, khớp hoàn toàn với 'alarming prevalence'."),
        ("[GIẢI THÍCH B]:", "SAI. Fleeting nghĩa là thoáng qua, mâu thuẫn với sự tồn tại lâu dài và tích tụ của vi nhựa."),
        ("[GIẢI THÍCH C]:", "SAI. Negligible nghĩa là không đáng kể, trái ngược với sắc thái 'alarming' (báo động)."),
        ("[GIẢI THÍCH D]:", "SAI. Extraneous nghĩa là xa lạ/không liên quan, không thể hiện quy mô phổ biến.")
    ]
    for tag, val in rationales_q1:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_tag(p, tag, val, tag_color="C2410C", val_color="334155")

    p_strat1 = doc.add_paragraph()
    p_strat1.paragraph_format.space_after = Pt(2)
    add_tag(p_strat1, "[MẸO GIẢI NHANH]:", "Chú ý từ manh mối 'alarming prevalence' ở câu trước để xác định sắc thái và quy mô của từ cần điền.", tag_color="4338CA", val_color="334155")

    p_mistake1 = doc.add_paragraph()
    p_mistake1.paragraph_format.space_after = Pt(4)
    add_tag(p_mistake1, "[CẢNH BÁO LỖI SAI]:", "Học sinh dễ nhầm 'extraneous' do nhầm lẫn nghĩa xa lạ với nghĩa rộng lớn.", tag_color="B91C1C", val_color="334155")

    # Q2 Solution
    add_styled_heading(doc, "Lời giải Câu 2 (Điền số)", level=2)
    doc_p2 = doc.add_paragraph()
    add_tag(doc_p2, "[CÂU 2]", tag_color="B91C1C")

    tags_q2 = [
        ("[ĐÁP ÁN ĐÚNG]:", "27 (hoặc 54/2)"),
        ("[THÔNG SỐ IRT]:", "a=1.50, b=1.10, c=0.00 (Phân hóa=1.5, Độ khó=1.1, Đoán mò=0)")
    ]
    for tag, val in tags_q2:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_tag(p, tag, val, tag_color="6B21A8", val_color="0369A1")

    p_exp2 = doc.add_paragraph()
    p_exp2.paragraph_format.space_before = Pt(4)
    p_exp2.paragraph_format.space_after = Pt(4)
    add_tag(p_exp2, "[LỜI GIẢI CHI TIẾT]:", tag_color="047857")
    r_exp2 = p_exp2.add_run("\n- Cách 1: Giải phương trình 3x + 9 = 21 => 3x = 12 => x = 4. Thế x = 4 vào 6x + 3 = 6(4) + 3 = 27.\n- Cách 2 (Nhanh): Biến đổi 6x + 3 = 2*(3x + 9) - 15 = 2(21) - 15 = 42 - 15 = 27.")
    r_exp2.font.name = 'Arial'

    p_strat2 = doc.add_paragraph()
    p_strat2.paragraph_format.space_after = Pt(2)
    add_tag(p_strat2, "[MẸO GIẢI NHANH]:", "So sánh hệ số của biến 6x và 3x để nhân trực tiếp cả biểu thức thay vì tìm x lẻ.", tag_color="4338CA", val_color="334155")

    p_mistake2 = doc.add_paragraph()
    p_mistake2.paragraph_format.space_after = Pt(4)
    add_tag(p_mistake2, "[CẢNH BÁO LỖI SAI]:", "Học sinh hay quên cộng 3 ở bước cuối sau khi tính ra 6x = 24.", tag_color="B91C1C", val_color="334155")

    doc.save(output_path)

if __name__ == "__main__":
    target_dir = r"c:\Users\minhc\Herd\digital-sat\docs\templates"
    os.makedirs(target_dir, exist_ok=True)
    
    exam_file = os.path.join(target_dir, "MAU_FILE_DE_THI_CHUAN_GIAO_VIEN.docx")
    sol_file = os.path.join(target_dir, "MAU_FILE_LOI_GIAI_CHUAN_GIAO_VIEN.docx")
    
    generate_exam_docx(exam_file)
    generate_solution_docx(sol_file)
    print(f"Successfully generated DOCX files with separate option tags in: {target_dir}")
